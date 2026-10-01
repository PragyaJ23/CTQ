"""Excel export of cohort matching RESULTS (not the input patients).

Two sheets (English workbook), plus two fully-Hindi sheets when lang=hi:
  * "Trial Results"  - one row per patient-trial pair: eligibility, match %,
                       reasons, failed criteria, criteria lists, locations.
  * "Patient Summary"- one row per patient: totals per verdict, best match,
                       and how the analysis was done (rule engine vs LLM).
  * "परिणाम (हिंदी)" / "मरीज़ सारांश (हिंदी)" - same data with Hindi headers,
                       verdicts, titles, reasons AND criteria (via the offline
                       English→Hindi glossary, services/en_hi.py - no LLM).

The frontend posts the same {patients, results} payload it already holds in
memory after a cohort run (structured OR unstructured - both use the same
cohort endpoint), so one endpoint serves both tabs. Result rows produced
before the Hindi fields existed fall back to translating title/reasons here.
"""
from io import BytesIO

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse

results_export_router = APIRouter(prefix="/api")

# Verdict labels in English and Hindi (order = ranking order).
EN_VERDICTS = ("Potentially Eligible", "Partially Eligible", "Not Eligible")
HI_VERDICTS = {
    "Potentially Eligible": "पात्र (Eligible)",
    "Partially Eligible": "आंशिक रूप से पात्र (Partially Eligible)",
    "Not Eligible": "पात्र नहीं (Not Eligible)",
}
HI_HEADERS = {
    "patient_id": "मरीज़ आईडी", "trial_id": "ट्रायल आईडी", "trial_title": "ट्रायल शीर्षक",
    "trial_condition": "ट्रायल की बीमारी", "phase": "चरण (Phase)", "status": "स्थिति",
    "sponsor": "प्रायोजक", "source": "स्रोत", "locations": "स्थान",
    "eligibility": "पात्रता", "match_percent": "मैच %", "similarity_score": "समानता स्कोर",
    "reasoning_method": "विश्लेषण विधि", "reasons_for": "पात्रता के कारण",
    "reasons_against": "अपात्रता के कारण", "missing_information": "अनुपलब्ध जानकारी",
    "failed_criteria": "असफल मानदंड",
    "inclusion_criteria": "समावेशन मानदंड", "exclusion_criteria": "अपवर्जन मानदंड",
    "age": "उम्र", "gender": "लिंग", "condition": "बीमारी",
    "trials_checked": "जाँचे गए ट्रायल", "potentially_eligible": "पूर्ण पात्र",
    "partially_eligible": "आंशिक पात्र", "not_eligible": "पात्र नहीं",
    "best_match_trial": "सर्वश्रेष्ठ मैच ट्रायल", "best_match_percent": "सर्वश्रेष्ठ मैच %",
    "best_match_title": "सर्वश्रेष्ठ मैच शीर्षक", "llm_used": "AI विश्लेषण",
}
HI_RESPONSES = {
    "Potentially Eligible": "पात्र",
    "Partially Eligible": "आंशिक रूप से पात्र",
    "Not Eligible": "पात्र नहीं",
}

_TRIAL_CACHE: dict | None = None


def _trial_by_id(trial_id):
    """Trial object for criteria lookup (local DB; never raises)."""
    global _TRIAL_CACHE
    if _TRIAL_CACHE is None:
        try:
            from services.trial_retrieval import load_trials

            _TRIAL_CACHE = {t.trial_id: t for t in load_trials()}
        except Exception:
            _TRIAL_CACHE = {}
    t = _TRIAL_CACHE.get(trial_id or "")
    if t is None:
        try:
            from services import database

            t = database.get_trial(trial_id or "")
        except Exception:
            t = None
    return t


def _hi(text) -> str:
    """Hindi rendering via the offline glossary (empty-safe)."""
    if not text:
        return ""
    from services.en_hi import to_hindi

    return to_hindi(text) or ""


def _hi_verdict(v) -> str:
    """Verdict in Hindi for sheet display."""
    return HI_RESPONSES.get(v, v or "")


def _s(v):
    """None-safe string for spreadsheet cells."""
    if v is None:
        return None
    return str(v) if not isinstance(v, str) else v


def _join(v):
    if not v:
        return None
    if isinstance(v, list):
        return "; ".join(_s(x) for x in v if x not in (None, "")) or None
    return _s(v)


def _cell(value):
    """Flatten a patient-input value into a spreadsheet-friendly cell."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, dict):
        return "; ".join(f"{k}: {v}" for k, v in value.items() if v not in (None, ""))
    if isinstance(value, list):
        if value and isinstance(value[0], dict):
            parts = []
            for m in value:
                bits = [str(m.get(k)) for k in ("name", "dose", "frequency", "duration_months")
                        if m.get(k) not in (None, "")]
                if bits:
                    parts.append(" ".join(bits))
            return "; ".join(parts)
        return "; ".join(str(v) for v in value if v not in (None, ""))
    return str(value)


@results_export_router.post("/cohort/export-results")
def export_results(payload: dict):
    """Download matching results as an Excel workbook.

    Payload: {"patients": [patient dicts], "results": [cohort result rows],
    "lang": "en" | "hi"} where each result row is {patient_id, ok, response}.
    """
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    patients = payload.get("patients") or []
    results = payload.get("results") or []
    if not results:
        raise HTTPException(status_code=400, detail="No results provided to export.")

    patient_rows = {(p or {}).get("patient_id"): p or {} for p in patients}

    wb = openpyxl.Workbook()

    # ------------------------------------------------------------------ #
    # Sheet 1: Trial Results - one row per patient-trial pair
    # ------------------------------------------------------------------ #
    ws = wb.active
    ws.title = "Trial Results"
    columns = [
        "patient_id", "trial_id", "trial_title", "trial_condition", "phase",
        "status", "sponsor", "source", "locations", "eligibility",
        "match_percent", "similarity_score", "reasoning_method",
        "reasons_for", "reasons_against", "missing_information", "failed_criteria",
        "inclusion_criteria", "exclusion_criteria",
    ]
    ws.append(columns)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    verdict_fills = {
        "Potentially Eligible": PatternFill("solid", fgColor="C6EFCE"),  # green
        "Partially Eligible": PatternFill("solid", fgColor="FFEB9C"),    # amber
        "Not Eligible": PatternFill("solid", fgColor="FFC7CE"),          # red
    }

    def _criteria_row(t: dict):
        """(inclusion, exclusion) joined strings from the trial database."""
        tr = _trial_by_id(t.get("trial_id"))
        inc = list(getattr(tr, "inclusion_criteria", []) or []) if tr else []
        exc = list(getattr(tr, "exclusion_criteria", []) or []) if tr else []
        return _join(inc), _join(exc)

    n_pairs = 0
    for r in results:
        pid = _s(r.get("patient_id"))
        if not r.get("ok"):
            ws.append([pid] + [""] * 8 + [f"ANALYSIS FAILED: {_s(r.get('error'))}"] + [None] * 9)
            n_pairs += 1
            continue
        for t in (r.get("response") or {}).get("results") or []:
            inc, exc = _criteria_row(t)
            ws.append([
                pid,
                _s(t.get("trial_id")),
                _s(t.get("title")),
                _s(t.get("condition")),
                _s(t.get("phase")),
                _s(t.get("status")),
                _s(t.get("sponsor")),
                _s(t.get("source")),
                _join(t.get("locations")),
                _s(t.get("eligibility")),
                t.get("match_percent"),
                t.get("similarity_score"),
                _s(t.get("reasoning_method")),
                _join(t.get("reasons_for")),
                _join(t.get("reasons_against")),
                _join(t.get("missing_information")),
                _join(t.get("failed_criteria")),
                inc,
                exc,
            ])
            n_pairs += 1
            row = ws[ws.max_row]
            fill = verdict_fills.get(t.get("eligibility"))
            if fill:
                row[9].fill = fill  # eligibility column
            row[9].alignment = Alignment(vertical="center", wrap_text=True)
            row[1].alignment = Alignment(vertical="center")

    widths = [12, 14, 42, 20, 8, 14, 18, 8, 34, 22, 13, 15, 14, 40, 40, 32, 40, 60, 60]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(columns))}{ws.max_row}"

    # ------------------------------------------------------------------ #
    # Sheet 2: Patient Summary - one row per patient
    # ------------------------------------------------------------------ #
    ws2 = wb.create_sheet("Patient Summary")
    sum_cols = ["patient_id", "age", "gender", "condition", "trials_checked",
                "potentially_eligible", "partially_eligible", "not_eligible",
                "best_match_trial", "best_match_percent", "best_match_title",
                "llm_used", "status"]
    ws2.append(sum_cols)
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    def _best_match(trials):
        counts = {v: 0 for v in EN_VERDICTS}
        for t in trials:
            counts[t.get("eligibility")] = counts.get(t.get("eligibility"), 0) + 1
        ranked = [t for t in trials if t.get("eligibility") == "Potentially Eligible"] \
            or sorted(trials, key=lambda t: (t.get("similarity_score") or 0), reverse=True)
        best = ranked[0] if ranked else None
        return counts, best

    for r in results:
        pid = _s(r.get("patient_id"))
        pin = patient_rows.get(pid, {})
        if not r.get("ok"):
            ws2.append([pid, _cell(pin.get("age")), _cell(pin.get("gender")),
                        _cell(pin.get("condition")), None, None, None, None,
                        None, None, None, None, f"FAILED: {_s(r.get('error'))}"])
            continue
        trials = (r.get("response") or {}).get("results") or []
        counts, best = _best_match(trials)
        ws2.append([
            pid, _cell(pin.get("age")), _cell(pin.get("gender")), _cell(pin.get("condition")),
            len(trials),
            counts.get("Potentially Eligible", 0),
            counts.get("Partially Eligible", 0),
            counts.get("Not Eligible", 0),
            best and _s(best.get("trial_id")),
            best and best.get("match_percent"),
            best and _s(best.get("title")),
            bool((r.get("response") or {}).get("llm_used")) or None,
            "ok",
        ])
        row = ws2[ws2.max_row]
        if counts.get("Potentially Eligible"):
            row[6].fill = verdict_fills["Potentially Eligible"]

    for idx, w in enumerate([12, 7, 9, 22, 13, 12, 11, 13, 14, 13, 42, 9, 30], start=1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = w
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(sum_cols))}{ws2.max_row}"

    # ------------------------------------------------------------------ #
    # Sheets 3+4 (lang=hi): FULLY Hindi - headers, verdicts, titles,
    # conditions, reasons AND inclusion/exclusion criteria (offline glossary).
    # ------------------------------------------------------------------ #
    if (payload.get("lang") or "").lower() in ("hi", "hindi", "हिंदी"):
        ws3 = wb.create_sheet("परिणाम (हिंदी)")
        ws3.append([HI_HEADERS.get(c, c) for c in columns])
        for cell in ws3[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")

        def _method_hi(v):
            return {"rule": "नियम-आधारित", "rule+llm": "नियम + AI"}.get(v, _s(v))

        for r in results:
            pid = _s(r.get("patient_id"))
            if not r.get("ok"):
                ws3.append([pid, "", "", "", "", "", "", "", "",
                            f"विश्लेषण असफल: {_s(r.get('error'))}"] + [None] * 9)
                continue
            for t in (r.get("response") or {}).get("results") or []:
                tr = _trial_by_id(t.get("trial_id"))
                inc_hi = [_hi(c) for c in (getattr(tr, "inclusion_criteria", None) or [])] if tr \
                    else [c for c in (t.get("inclusion_criteria") or [])]
                exc_hi = [_hi(c) for c in (getattr(tr, "exclusion_criteria", None) or [])] if tr \
                    else [c for c in (t.get("exclusion_criteria") or [])]
                ws3.append([
                    pid,
                    _s(t.get("trial_id")),
                    t.get("title_hi") or _hi(t.get("title")),
                    t.get("condition_hi") or _hi(t.get("condition")),
                    _s(t.get("phase")),
                    _s(t.get("status")),
                    _s(t.get("sponsor")),
                    _s(t.get("source")),
                    _join(t.get("locations")),
                    _hi_verdict(t.get("eligibility")),
                    t.get("match_percent"),
                    t.get("similarity_score"),
                    _method_hi(t.get("reasoning_method")),
                    _join(t.get("reasons_for_hi") or [_hi(x) for x in (t.get("reasons_for") or [])]),
                    _join(t.get("reasons_against_hi") or [_hi(x) for x in (t.get("reasons_against") or [])]),
                    _join(t.get("missing_information_hi") or [_hi(x) for x in (t.get("missing_information") or [])]),
                    _join(t.get("failed_criteria_hi") or [_hi(x) for x in (t.get("failed_criteria") or [])]),
                    _join(inc_hi) or None,
                    _join(exc_hi) or None,
                ])
                row = ws3[ws3.max_row]
                fill = verdict_fills.get(t.get("eligibility"))
                if fill:
                    row[9].fill = fill

        for idx, w in enumerate(widths, start=1):
            ws3.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = w
        ws3.freeze_panes = "A2"
        ws3.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(columns))}{ws3.max_row}"

        ws4 = wb.create_sheet("मरीज़ सारांश (हिंदी)")
        ws4.append([HI_HEADERS.get(c, c) for c in sum_cols])
        for cell in ws4[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for r in results:
            pid = _s(r.get("patient_id"))
            pin = patient_rows.get(pid, {})
            if not r.get("ok"):
                ws4.append([pid, _cell(pin.get("age")), _cell(pin.get("gender")),
                            _cell(pin.get("condition")), None, None, None, None,
                            None, None, None, None, f"असफल: {_s(r.get('error'))}"])
                continue
            trials = (r.get("response") or {}).get("results") or []
            counts, best = _best_match(trials)
            best_title = None
            if best:
                best_title = best.get("title_hi") or _hi(best.get("title")) or _s(best.get("title"))
            ws4.append([
                pid, _cell(pin.get("age")), _cell(pin.get("gender")),
                _hi(pin.get("condition")) or _cell(pin.get("condition")),
                len(trials),
                counts.get("Potentially Eligible", 0),
                counts.get("Partially Eligible", 0),
                counts.get("Not Eligible", 0),
                best and _s(best.get("trial_id")),
                best and best.get("match_percent"),
                best_title,
                bool((r.get("response") or {}).get("llm_used")) or None,
                "पूर्ण",
            ])
        for idx, w in enumerate([12, 7, 9, 22, 13, 12, 11, 13, 14, 13, 42, 9, 30], start=1):
            ws4.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = w
        ws4.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="ctq_matching_results_hindi.xlsx"'
                 if (payload.get("lang") or "").lower() in ("hi", "hindi", "हिंदी")
                 else 'attachment; filename="ctq_matching_results.xlsx"'},
    )
